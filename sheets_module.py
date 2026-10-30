"""Combine assignment data and write to a local Excel assessment tracker."""

import logging
import re
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from config import EXCEL_FILE_PATH
from course_utils import prefer_course_label
from vocab import prefer_status
from date_utils import (
    apply_lab_monday_due,
    has_explicit_time,
    prefer_timed_due,
    split_title_and_due,
)
from dedupe_module import (
    TASK_ID_COL,
    dedupe_cross_source,
    fuzzy_match_key,
    identity_key,
    is_droppable_placeholder,
    remember_task_ids,
    records_to_tracker_frame,
)

logger = logging.getLogger(__name__)

SHEET_NAME = "Assessment Schedule"
AUTO_COLUMNS = ["Course", "Assessment title", "Due Date", TASK_ID_COL, "Is Draft"]
MANUAL_COLUMNS = [
    "Estimated time dedicated to task",
    "Status",
    "Priority",
]
# Task_ID is tracking-only; written to Excel but not required for display polish
COLUMNS = [
    "Course",
    "Assessment title",
    "Due Date",
    "Estimated time dedicated to task",
    "Status",
    "Priority",
    TASK_ID_COL,
    "Is Draft",
]

TITLE = "Assessment Schedule"
TITLE_FONT = Font(name="Calibri", bold=True, size=16, color="FFFFFF")
TITLE_FILL = PatternFill("solid", fgColor="2F5496")

HEADER_FONT = Font(name="Calibri", bold=True, size=11, color="1F3864")
HEADER_FILL = PatternFill("solid", fgColor="FBF3CB")
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)

BODY_FONT = Font(name="Calibri", size=11, color="333333")
ALT_ROW_FILL = PatternFill("solid", fgColor="F5F7FA")
WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")

THIN_BORDER = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)

COLUMN_WIDTHS = {
    "A": 14.0,
    "B": 28.0,
    "C": 16.0,
    "D": 22.0,
    "E": 14.0,
    "F": 12.0,
    "G": 28.0,  # Task_ID (tracking)
    "H": 10.0,  # Is Draft
}

COLUMN_ALIGNMENTS = {
    "Course": Alignment(horizontal="left", vertical="center"),
    "Assessment title": Alignment(horizontal="left", vertical="center", wrap_text=True),
    "Due Date": Alignment(horizontal="center", vertical="center"),
    "Estimated time dedicated to task": Alignment(horizontal="center", vertical="center"),
    "Status": Alignment(horizontal="center", vertical="center"),
    "Priority": Alignment(horizontal="center", vertical="center"),
    TASK_ID_COL: Alignment(horizontal="left", vertical="center"),
    "Is Draft": Alignment(horizontal="center", vertical="center"),
}

STATUS_OPTIONS = '"Not started,In Progress,Done"'
PRIORITY_OPTIONS = '"P 1,P 2,P 3"'


def _normalize_key(course: Any, title: Any) -> tuple[str, str]:
    """
    Fuzzy match key so near-duplicate titles collapse.

    Returns a 2-tuple for backward compatibility with google_sheets_module.
    """
    key = fuzzy_match_key(course, title)
    course_part, _, title_part = key.partition("||")
    return course_part, title_part


def _prefer_title(current: str, candidate: str) -> str:
    """Keep the cleaner, more descriptive title when collapsing duplicates."""
    cur = str(current or "").strip()
    cand = str(candidate or "").strip()
    if not cur:
        return cand
    if not cand:
        return cur
    # Prefer titles that do not still embed "due date …"
    cur_has_due = bool(re.search(r"\bdue\b", cur, re.IGNORECASE))
    cand_has_due = bool(re.search(r"\bdue\b", cand, re.IGNORECASE))
    if cur_has_due and not cand_has_due:
        return cand
    if cand_has_due and not cur_has_due:
        return cur
    if len(cand) > len(cur):
        return cand
    return cur


def _prefer_task_id(current: Any, candidate: Any) -> str:
    """Prefer Canvas-backed Task_IDs over syllabus composites."""
    cur = str(current or "").strip()
    cand = str(candidate or "").strip()
    if cand.startswith("canvas_") and not cur.startswith("canvas_"):
        return cand
    if cur:
        return cur
    return cand


def _collapse_by_match_key(df: pd.DataFrame) -> pd.DataFrame:
    """Collapse rows that share the same fuzzy (course, title) match key."""
    cols = [c for c in COLUMNS if c in df.columns] or list(df.columns)
    if df.empty:
        return df.reindex(columns=cols)

    best: dict[str, dict[str, Any]] = {}
    # Canvas rows group by their Canvas Task_ID (two Canvas items never merge);
    # non-Canvas rows join the first Canvas row sharing their fuzzy key, else
    # group by fuzzy key as before.
    fuzzy_to_canvas: dict[str, str] = {}
    if TASK_ID_COL in df.columns:
        for _, r in df.iterrows():
            tid = str(r.get(TASK_ID_COL, "") or "").strip()
            if tid.startswith("canvas_"):
                fuzzy_to_canvas.setdefault(
                    fuzzy_match_key(r.get("Course", ""), r.get("Assessment title", "")), tid
                )
    for _, row in df.iterrows():
        key = fuzzy_match_key(row.get("Course", ""), row.get("Assessment title", ""))
        if key == "||":
            continue
        ikey = identity_key(
            row.get("Course", ""), row.get("Assessment title", ""), row.get(TASK_ID_COL, "")
        )
        key = ikey if ikey != key else fuzzy_to_canvas.get(key, key)

        if key not in best:
            best[key] = {col: row.get(col, "") for col in df.columns}
            # Canonicalize course label on first insert
            if "Course" in best[key]:
                best[key]["Course"] = prefer_course_label(
                    "", str(best[key].get("Course", ""))
                )
            continue

        existing = best[key]
        existing["Assessment title"] = _prefer_title(
            str(existing.get("Assessment title", "")),
            str(row.get("Assessment title", "")),
        )
        if "Course" in df.columns:
            existing["Course"] = prefer_course_label(
                str(existing.get("Course", "")),
                str(row.get("Course", "")),
            )
        if TASK_ID_COL in df.columns:
            existing[TASK_ID_COL] = _prefer_task_id(
                existing.get(TASK_ID_COL, ""), row.get(TASK_ID_COL, "")
            )
        if "Due Date" in df.columns or "Due Date" in existing:
            existing["Due Date"] = prefer_timed_due(
                existing.get("Due Date", ""), row.get("Due Date", "")
            )
        for col in MANUAL_COLUMNS:
            if col == "Status" and col in df.columns:
                existing[col] = prefer_status(existing.get(col, ""), row.get(col, ""))
            elif col in df.columns and not str(existing.get(col, "")).strip():
                existing[col] = row.get(col, "")
        # Same COURSE 202 lab quiz from two labels/dates -> one item on the Monday.
        if "Due Date" in existing:
            existing["Due Date"] = apply_lab_monday_due(
                existing.get("Course", ""),
                existing.get("Assessment title", ""),
                existing.get("Due Date", ""),
            )

    if not best:
        return pd.DataFrame(columns=cols)

    return pd.DataFrame(list(best.values()), columns=list(df.columns))


IGNORED_ITEMS_PATH = Path(__file__).resolve().parent / "ignored_items.json"


def load_ignored_task_ids(path: Path | None = None) -> set[str]:
    """Task_IDs listed in ignored_items.json ({"ignored": [{"task_id": ...}, ...]})."""
    import json

    p = Path(path) if path else IGNORED_ITEMS_PATH
    if not p.exists():
        return set()
    try:
        data = json.loads(p.read_text(encoding="utf-8-sig"))
    except (OSError, json.JSONDecodeError) as exc:
        logger.warning("Could not read %s: %s", p, exc)
        return set()
    items = data.get("ignored", []) if isinstance(data, dict) else []
    out: set[str] = set()
    for item in items:
        tid = item.get("task_id") if isinstance(item, dict) else item
        if str(tid or "").strip():
            out.add(str(tid).strip())
    return out


def _drop_ignored_items(frame: pd.DataFrame) -> pd.DataFrame:
    """Explicitly suppressed items (by Task_ID) are never re-added from sources."""
    ignored = load_ignored_task_ids()
    if not ignored or frame.empty or TASK_ID_COL not in frame.columns:
        return frame
    mask = frame[TASK_ID_COL].astype(str).str.strip().isin(ignored)
    if mask.any():
        logger.info(
            "Suppressed %d item(s) listed in ignored_items.json: %s",
            int(mask.sum()),
            sorted(frame.loc[mask, TASK_ID_COL].astype(str)),
        )
    return frame[~mask].reset_index(drop=True)


def _restrict_to_allowed_courses(frame: pd.DataFrame) -> pd.DataFrame:
    """Incoming rows outside ALLOWED_COURSES never reach the tracker (existing rows are untouched)."""
    from config import ALLOWED_COURSES
    from course_utils import course_matches_allowed, parse_allowed_courses

    allowed = parse_allowed_courses(ALLOWED_COURSES)
    if not allowed or frame.empty or "Course" not in frame.columns:
        return frame
    mask = frame["Course"].astype(str).map(lambda c: course_matches_allowed(c, allowed))
    if (~mask).any():
        logging.getLogger(__name__).warning(
            "Dropped %d incoming row(s) outside ALLOWED_COURSES: %s",
            int((~mask).sum()),
            sorted(set(frame.loc[~mask, "Course"].astype(str))),
        )
    return frame[mask].reset_index(drop=True)


def _incoming_from_sources(
    canvas_data: list[dict[str, Any]],
    syllabus_data: list[dict[str, Any]],
):
    """
    Dedupe Canvas + syllabus, then map to tracker auto-fill columns.

    Returns:
        (incoming DataFrame with Task_ID, DedupStats)
    """
    records, stats = dedupe_cross_source(canvas_data, syllabus_data)
    frame = _drop_ignored_items(_restrict_to_allowed_courses(records_to_tracker_frame(records)))
    cols = AUTO_COLUMNS + ["Source"]
    return frame.reindex(columns=cols).fillna(""), stats


def _load_existing_tracker(path: Path) -> pd.DataFrame:
    """Load an existing tracker file, or return an empty frame."""
    if not path.exists():
        return pd.DataFrame(columns=COLUMNS)

    for header_row in (1, 0):
        try:
            df = pd.read_excel(
                path, sheet_name=SHEET_NAME, engine="openpyxl", header=header_row
            )
        except ValueError:
            df = pd.read_excel(path, engine="openpyxl", header=header_row)

        if "Course" in df.columns:
            df = df.reindex(columns=COLUMNS)
            return df.fillna("")

    return pd.DataFrame(columns=COLUMNS)


def _sort_by_days_until_due(df: pd.DataFrame) -> pd.DataFrame:
    """
    Sort the tracker so "what comes next" is at the top.

    Order:
      1. Active work (Not Started / In Progress / …) before finished
         (Submitted / Complete / Graded / Cancelled) — finished sinks to bottom
      2. Soonest due date first (past-due active work stays near the top)
      3. Earlier time-of-day within the same calendar day
      4. Course, then title
      5. Missing / unparseable due dates sink to the bottom of their group
    """
    if df.empty or "Due Date" not in df.columns:
        return df.reset_index(drop=True)

    from vocab import Status, normalize_status

    finished = {
        Status.COMPLETE,
        Status.SUBMITTED,
        Status.GRADED,
        Status.CANCELLED,
    }

    out = df.copy()
    out["_sort_date"] = pd.to_datetime(out["Due Date"], errors="coerce")
    out["_sort_day"] = out["_sort_date"].dt.normalize()
    if "Status" in out.columns:
        out["_done"] = [
            1 if normalize_status(s) in finished else 0 for s in out["Status"].tolist()
        ]
    else:
        out["_done"] = 0

    out = out.sort_values(
        by=["_done", "_sort_day", "_sort_date", "Course", "Assessment title"],
        ascending=[True, True, True, True, True],
        na_position="last",
    )
    return out.drop(columns=["_sort_date", "_sort_day", "_done"]).reset_index(drop=True)


def _is_droppable_placeholder(title: Any, due: Any = "") -> bool:
    """Bare syllabus headings that must never remain on the tracker."""
    return is_droppable_placeholder(title, due)


DUE_OVERRIDES_PATH = Path(__file__).resolve().parent / "due_overrides.json"


def load_due_overrides(path: Path | None = None) -> dict[str, dict[str, Any]]:
    """
    Read manual due/end-time overrides keyed by fuzzy (course, title).

    File format: {"overrides": [{"course": "COURSE 401", "title": "Midterm",
    "due": "2026-10-23 11:00", "end": "2026-10-23 11:50"}]}. Missing file -> {}.
    """
    import json

    p = Path(path) if path else DUE_OVERRIDES_PATH
    if not p.exists():
        return {}
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        logger.warning("Could not read %s: %s", p, exc)
        return {}
    out: dict[str, dict[str, Any]] = {}
    items = data.get("overrides", []) if isinstance(data, dict) else []
    for item in items:
        if not isinstance(item, dict) or not str(item.get("due") or "").strip():
            continue
        key = fuzzy_match_key(item.get("course", ""), item.get("title", ""))
        if key != "||":
            out[key] = item
    return out


def apply_due_overrides(
    df: pd.DataFrame, overrides: dict[str, dict[str, Any]] | None = None
) -> pd.DataFrame:
    """Manual due/end overrides win over any source (e.g. date-only announcements)."""
    overrides = load_due_overrides() if overrides is None else overrides
    if df.empty or not overrides:
        return df
    out = df.copy()
    if any(str(o.get("end") or "").strip() for o in overrides.values()) and "End Date" not in out.columns:
        out["End Date"] = ""
    for idx, row in out.iterrows():
        ov = overrides.get(fuzzy_match_key(row.get("Course", ""), row.get("Assessment title", "")))
        if not ov:
            continue
        out.at[idx, "Due Date"] = str(ov["due"]).strip()
        end = str(ov.get("end") or "").strip()
        if end:
            out.at[idx, "End Date"] = end
    return out


def _merge_tracker(existing: pd.DataFrame, incoming: pd.DataFrame) -> pd.DataFrame:
    """
    Merge synced data with the existing tracker.

    - Matching rows (Task_ID or fuzzy title key): update auto columns, keep manual columns.
    - New rows: add with blank manual columns.
    - Rows only in existing: keep (preserves completed/historical entries),
      except bare category phantoms like "Labs" / "Assignments".
    - Near-duplicate titles within the same course are collapsed.
    """
    existing = existing.reindex(columns=COLUMNS).fillna("")
    extra_cols = [c for c in ("Source", "End Date") if c in incoming.columns]
    incoming = incoming.reindex(
        columns=[c for c in AUTO_COLUMNS + extra_cols if c in incoming.columns or c in AUTO_COLUMNS]
    ).fillna("")
    for col in AUTO_COLUMNS:
        if col not in incoming.columns:
            incoming[col] = ""

    # Clean embedded dues + drop category placeholders before matching
    def _sanitize_frame(df: pd.DataFrame) -> pd.DataFrame:
        if df.empty:
            return df
        rows: list[dict[str, Any]] = []
        for _, row in df.iterrows():
            title = str(row.get("Assessment title", "") or "")
            due = row.get("Due Date", "")
            title, due = split_title_and_due(title, due)
            due = apply_lab_monday_due(row.get("Course", ""), title, due)
            if _is_droppable_placeholder(title, due):
                continue
            rec = {col: row.get(col, "") for col in df.columns}
            rec["Assessment title"] = title
            rec["Due Date"] = due
            rows.append(rec)
        return pd.DataFrame(rows, columns=list(df.columns)) if rows else pd.DataFrame(columns=df.columns)

    existing = _sanitize_frame(existing)
    incoming = _sanitize_frame(incoming)

    existing = _collapse_by_match_key(existing)
    incoming = _collapse_by_match_key(incoming)

    manual_by_key: dict[str, dict[str, Any]] = {}
    manual_by_task_id: dict[str, dict[str, Any]] = {}
    existing_title_by_key: dict[str, str] = {}
    existing_task_id_by_key: dict[str, str] = {}
    existing_due_by_key: dict[str, Any] = {}

    fuzzy_to_existing: dict[str, str] = {}
    fuzzy_to_existing_canvas: dict[str, str] = {}
    for _, row in existing.iterrows():
        fkey = fuzzy_match_key(row["Course"], row["Assessment title"])
        if fkey == "||":
            continue
        key = identity_key(row["Course"], row["Assessment title"], row.get(TASK_ID_COL, ""))
        if key == fkey:
            fuzzy_to_existing.setdefault(fkey, key)
        else:
            fuzzy_to_existing_canvas.setdefault(fkey, key)
        manual = {col: row.get(col, "") for col in MANUAL_COLUMNS}
        manual_by_key[key] = manual
        existing_title_by_key[key] = str(row.get("Assessment title", ""))
        existing_due_by_key[key] = row.get("Due Date", "")
        task_id = str(row.get(TASK_ID_COL, "") or "").strip()
        if task_id:
            manual_by_task_id[task_id] = manual
            existing_task_id_by_key[key] = task_id

    merged_rows: list[dict[str, Any]] = []
    incoming_keys: set[str] = set()
    incoming_task_ids: set[str] = set()

    incoming_fuzzy: set[str] = set()
    for _, row in incoming.iterrows():
        fkey = fuzzy_match_key(row["Course"], row["Assessment title"])
        if fkey == "||":
            continue
        own_tid = str(row.get(TASK_ID_COL, "") or "").strip()
        ikey = identity_key(row["Course"], row["Assessment title"], own_tid)
        if ikey in manual_by_key:
            key = ikey
        elif ikey != fkey:
            # Canvas item not on the sheet yet: it may adopt a syllabus row
            # (promotion) but never a different Canvas item with a similar title.
            key = fuzzy_to_existing.get(fkey, ikey)
        else:
            key = fuzzy_to_existing.get(fkey) or fuzzy_to_existing_canvas.get(fkey) or ikey
        incoming_fuzzy.add(fkey)

        task_id = own_tid or existing_task_id_by_key.get(key, "")
        incoming_keys.add(key)
        if task_id:
            incoming_task_ids.add(task_id)

        by_tid = manual_by_task_id.get(task_id) or {}
        by_key = manual_by_key.get(key) or {}
        manual = dict(by_tid or by_key or {col: "" for col in MANUAL_COLUMNS})
        if by_tid and by_key:
            # Task_ID and fuzzy-key rows can differ; never downgrade the status.
            manual["Status"] = prefer_status(by_tid.get("Status", ""), by_key.get("Status", ""))
        title = _prefer_title(
            existing_title_by_key.get(key, ""),
            str(row.get("Assessment title", "")),
        )
        # Incoming wins when both are timed or both are date-only. A date-only
        # or blank incoming value must not wipe a clock time already stored.
        incoming_due = row.get("Due Date", "")
        existing_due = existing_due_by_key.get(key, "")
        if has_explicit_time(existing_due) and not has_explicit_time(incoming_due):
            due = existing_due
        else:
            due = incoming_due
        title, due = split_title_and_due(title, due)
        due = apply_lab_monday_due(row["Course"], title, due)
        if _is_droppable_placeholder(title, due):
            continue
        merged = {
            "Course": row["Course"],
            "Assessment title": title,
            "Due Date": due,
            TASK_ID_COL: task_id,
            "Is Draft": row.get("Is Draft", False),
            **manual,
        }
        end_date = str(row.get("End Date", "") or "").strip()
        if end_date:
            merged["End Date"] = end_date
        source = str(row.get("Source", "") or "").strip()
        if source:
            merged["Source"] = source
        merged_rows.append(merged)

    for _, row in existing.iterrows():
        key = fuzzy_match_key(row["Course"], row["Assessment title"])
        if key == "||":
            continue
        task_id = str(row.get(TASK_ID_COL, "") or "").strip()
        ikey = identity_key(row["Course"], row["Assessment title"], task_id)
        if ikey in incoming_keys or (task_id and task_id in incoming_task_ids):
            continue
        if ikey == key and key in incoming_fuzzy:
            # Non-Canvas row superseded by an incoming item with the same fuzzy title
            continue
        title = str(row.get("Assessment title", "") or "")
        due = row.get("Due Date", "")
        title, due = split_title_and_due(title, due)
        due = apply_lab_monday_due(row.get("Course", ""), title, due)
        # Drop stale phantoms that were written before we filtered them
        if _is_droppable_placeholder(title, due):
            continue
        kept = {col: row.get(col, "") for col in COLUMNS}
        kept["Assessment title"] = title
        kept["Due Date"] = due
        merged_rows.append(kept)

    out_cols = list(COLUMNS)
    for extra in ("Source", "End Date"):
        if any(extra in row for row in merged_rows) and extra not in out_cols:
            out_cols.append(extra)

    if not merged_rows:
        return pd.DataFrame(columns=out_cols)

    df = pd.DataFrame(merged_rows)
    for col in out_cols:
        if col not in df.columns:
            df[col] = ""
    df = df.reindex(columns=out_cols).fillna("")
    df = _collapse_by_match_key(df)
    # Final pass after title cleanup — catches Online Assignment 1 vs …
    df = _sanitize_frame(df)
    df = _collapse_by_match_key(df)
    df = apply_due_overrides(df)
    df = _sort_by_days_until_due(df)
    remember_task_ids(df)
    return df


def _apply_title_row(ws) -> int:
    """Add a merged title banner. Returns the header row index."""
    last_col = get_column_letter(len(COLUMNS))
    ws.merge_cells(f"A1:{last_col}1")
    title_cell = ws["A1"]
    title_cell.value = TITLE
    title_cell.font = TITLE_FONT
    title_cell.fill = TITLE_FILL
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 34
    return 2


def _apply_header_row(ws, header_row: int) -> None:
    """Style the column header row."""
    ws.row_dimensions[header_row].height = 36
    for col_idx, header in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER


def _apply_body_styles(ws, header_row: int, data_row_count: int) -> None:
    """Apply zebra striping, borders, and alignment to data rows."""
    first_data_row = header_row + 1
    last_data_row = header_row + data_row_count

    for row_idx in range(first_data_row, last_data_row + 1):
        is_alt = (row_idx - first_data_row) % 2 == 1
        row_fill = ALT_ROW_FILL if is_alt else WHITE_FILL
        ws.row_dimensions[row_idx].height = 22

        for col_idx, column in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = BODY_FONT
            cell.fill = row_fill
            cell.border = THIN_BORDER
            cell.alignment = COLUMN_ALIGNMENTS.get(
                column, Alignment(vertical="center")
            )

            if column == "Due Date" and cell.value:
                cell.number_format = "ddd, mmm d, yyyy"


def _apply_validations(ws, header_row: int, data_row_count: int) -> None:
    """Add dropdown lists for Status and Priority."""
    if data_row_count == 0:
        return

    first_row = header_row + 1
    last_row = max(header_row + data_row_count, header_row + 50)

    status_col = get_column_letter(COLUMNS.index("Status") + 1)
    priority_col = get_column_letter(COLUMNS.index("Priority") + 1)

    status_dv = DataValidation(type="list", formula1=STATUS_OPTIONS, allow_blank=True)
    status_dv.error = "Pick a status from the list."
    status_dv.prompt = "Select status"
    ws.add_data_validation(status_dv)
    status_dv.add(f"{status_col}{first_row}:{status_col}{last_row}")

    priority_dv = DataValidation(type="list", formula1=PRIORITY_OPTIONS, allow_blank=True)
    priority_dv.error = "Pick a priority from the list."
    priority_dv.prompt = "Select priority"
    ws.add_data_validation(priority_dv)
    priority_dv.add(f"{priority_col}{first_row}:{priority_col}{last_row}")


def _apply_conditional_formatting(ws, header_row: int, data_row_count: int) -> None:
    """Highlight overdue dates and color-code status values."""
    if data_row_count == 0:
        return

    first_row = header_row + 1
    last_row = header_row + data_row_count
    due_col = get_column_letter(COLUMNS.index("Due Date") + 1)
    status_col = get_column_letter(COLUMNS.index("Status") + 1)

    overdue_fill = PatternFill("solid", fgColor="FCE4D6")
    overdue_font = Font(name="Calibri", size=11, color="9C0006", bold=True)
    ws.conditional_formatting.add(
        f"{due_col}{first_row}:{due_col}{last_row}",
        FormulaRule(
            formula=[f'AND(${due_col}{first_row}<>"",${due_col}{first_row}<TODAY())'],
            fill=overdue_fill,
            font=overdue_font,
        ),
    )

    done_fill = PatternFill("solid", fgColor="E2EFDA")
    done_font = Font(name="Calibri", size=11, color="375623")
    ws.conditional_formatting.add(
        f"A{first_row}:{get_column_letter(len(COLUMNS))}{last_row}",
        FormulaRule(
            formula=[f'${status_col}{first_row}="Done"'],
            fill=done_fill,
            font=done_font,
        ),
    )


def _write_tracker(path: Path, df: pd.DataFrame) -> None:
    """Write a styled assessment tracker workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.sheet_view.showGridLines = False

    header_row = _apply_title_row(ws)
    _apply_header_row(ws, header_row)

    data_start_row = header_row + 1
    for row_idx, record in enumerate(df.to_dict(orient="records"), start=data_start_row):
        for col_idx, column in enumerate(COLUMNS, start=1):
            value = record.get(column, "")
            if column == "Due Date" and value:
                parsed = pd.to_datetime(value, errors="coerce")
                value = parsed.to_pydatetime() if pd.notna(parsed) else value
            ws.cell(row=row_idx, column=col_idx, value=value or None)

    data_row_count = len(df)
    _apply_body_styles(ws, header_row, data_row_count)
    _apply_validations(ws, header_row, data_row_count)
    _apply_conditional_formatting(ws, header_row, data_row_count)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1).coordinate
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(COLUMNS))}"
        f"{max(header_row + data_row_count, header_row)}"
    )

    for letter, width in COLUMN_WIDTHS.items():
        ws.column_dimensions[letter].width = width

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def update_sheet(
    canvas_data: list[dict[str, Any]],
    syllabus_data: list[dict[str, Any]],
) -> pd.DataFrame:
    """
    Combine data sources, merge with existing tracker, and write Excel file.

    Returns:
        The combined DataFrame that was written to the file.
    """
    output_path = Path(EXCEL_FILE_PATH)
    existing = _load_existing_tracker(output_path)
    incoming, stats = _incoming_from_sources(canvas_data, syllabus_data)
    from dedupe_module import print_dedup_stats

    print_dedup_stats(stats)

    if incoming.empty and existing.empty:
        logger.warning("No assignment data to write.")

    df = _merge_tracker(existing, incoming)

    try:
        _write_tracker(output_path, df)
        logger.info("Wrote %d row(s) to %s.", len(df), output_path.resolve())
    except OSError as exc:
        logger.error("Failed to write Excel file: %s", exc)
        raise

    return df


# --- Public API (stable imports for other modules) ---
normalize_key = _normalize_key
sort_by_days_until_due = _sort_by_days_until_due
incoming_from_sources = _incoming_from_sources
merge_tracker = _merge_tracker
